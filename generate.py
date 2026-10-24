#!/usr/bin/env python3
"""IQKnacker Video-Generator.

Beispiele:
    python generate.py                 # alle Rätsel mit Status "neu"
    python generate.py --id 5          # nur Rätsel 5 (egal welcher Status)
    python generate.py --id 3 7 9      # mehrere Rätsel
    python generate.py --alle          # alle Rätsel neu erzeugen
    python generate.py --vorschau      # nur Vorschaubilder (PNG), kein Video
    python generate.py --liste         # Rätsel mit Status anzeigen
    python generate.py --check         # Installation prüfen
    python generate.py --neu-emoji 5   # 5 neue Emoji-Gleichungen anlegen (Lösung rechnet das Skript)
    python generate.py --neu-finde 5   # 5 neue „Finde die 9“-Rätsel anlegen
    python generate.py --plan 7        # Posting-Plan für 7 Tage (3 Videos/Tag) in output/postingplan.txt
"""

import argparse
import array
import json
import math
import os
import random
import shutil
import subprocess
import sys
import datetime
import time
import wave

BASIS = os.path.dirname(os.path.abspath(__file__))
CONFIG_DATEI = os.path.join(BASIS, "config.json")
RAETSEL_DATEI = os.path.join(BASIS, "raetsel.json")
PFLICHTFELDER = ["id", "kategorie", "frage", "loesung", "erklaerung", "schwierigkeit", "status"]
GUELTIGE_STATUS = ["neu", "erstellt", "gepostet"]
MUSIK_ENDUNGEN = (".mp3", ".m4a", ".aac", ".wav", ".ogg")
KOMMENTAR_MAX = 150  # Zeichenlimit für TikTok-Kommentare


# ---------------------------------------------------------------- Umgebung prüfen

def pip_installieren():
    print("→ Installiere fehlende Python-Pakete aus requirements.txt ...")
    befehl = [sys.executable, "-m", "pip", "install", "-r", os.path.join(BASIS, "requirements.txt")]
    if subprocess.call(befehl) != 0:
        sys.exit("✗ Installation fehlgeschlagen. Bitte manuell ausführen:\n  " + " ".join(befehl))


def pakete_sicherstellen():
    try:
        import numpy  # noqa: F401
        import PIL  # noqa: F401
    except ImportError:
        pip_installieren()


def ffmpeg_finden(installieren=True):
    pfad = shutil.which("ffmpeg")
    if pfad:
        return pfad
    try:
        import imageio_ffmpeg
        return imageio_ffmpeg.get_ffmpeg_exe()
    except Exception:
        pass
    if installieren:
        pip_installieren()
        return ffmpeg_finden(installieren=False)
    sys.exit("✗ FFmpeg nicht gefunden. Installiere es z. B. mit\n"
             "  macOS:   brew install ffmpeg\n"
             "  Windows: winget install ffmpeg\n"
             "  Linux:   sudo apt install ffmpeg")


def check():
    print(f"Python:  {sys.version.split()[0]}", "✓" if sys.version_info >= (3, 8) else "✗ (mind. 3.8 nötig)")
    pakete_sicherstellen()
    import numpy
    import PIL
    print(f"Pillow:  {PIL.__version__} ✓")
    print(f"numpy:   {numpy.__version__} ✓")
    ffmpeg = ffmpeg_finden()
    version = subprocess.run([ffmpeg, "-version"], capture_output=True, text=True).stdout.split("\n")[0]
    print(f"FFmpeg:  {version} ✓\n         ({ffmpeg})")
    cfg = lade_json(CONFIG_DATEI)
    for name, pfad in cfg["schriften"].items():
        voll = os.path.join(BASIS, pfad)
        print(f"Schrift {name}: {pfad}", "✓" if os.path.exists(voll) else "✗ fehlt")
    musik = musik_dateien(cfg)
    modus = cfg["musik"].get("modus", "auto")
    if modus == "aus":
        info = "aus"
    elif musik and modus != "generiert":
        info = f"{len(musik)} Track(s) aus {cfg['musik']['ordner']}"
    elif modus in ("auto", "generiert"):
        info = "wird automatisch erzeugt (keine Dateien nötig)"
    else:
        info = f"keine Tracks in {cfg['musik']['ordner']} – Videos ohne Musik"
    print(f"Musik:   {info}")
    logo = os.path.join(BASIS, cfg["wasserzeichen"]["datei"])
    print(f"Logo:    {'✓ ' + cfg['wasserzeichen']['datei'] if os.path.exists(logo) else '– nicht gefunden, nutze Text-Wasserzeichen'}")


# ---------------------------------------------------------------- Daten

def lade_json(pfad):
    with open(pfad, encoding="utf-8") as f:
        return json.load(f)


def speichere_raetsel(raetsel):
    tmp = RAETSEL_DATEI + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(raetsel, f, ensure_ascii=False, indent=2)
        f.write("\n")
    os.replace(tmp, RAETSEL_DATEI)


def termine_abhaken(raetsel, jetzt=None):
    """Rätsel, deren geplanter Termin („termin“ aus --plan) vorbei ist, gelten als gepostet."""
    jetzt = jetzt or datetime.datetime.now()
    erledigt = [r for r in raetsel if r.get("termin") and r["status"] != "gepostet"
                and datetime.datetime.strptime(r["termin"], "%Y-%m-%d %H:%M") <= jetzt]
    for r in erledigt:
        r["status"] = "gepostet"
    return erledigt


def pruefe_raetsel(raetsel, cfg):
    fehler, ids = [], set()
    for i, r in enumerate(raetsel):
        fehlend = [f for f in PFLICHTFELDER if f not in r]
        if fehlend:
            fehler.append(f"Eintrag {i + 1}: Felder fehlen: {', '.join(fehlend)}")
            continue
        if r["id"] in ids:
            fehler.append(f"Rätsel-ID {r['id']} kommt doppelt vor")
        ids.add(r["id"])
        if r["status"] not in GUELTIGE_STATUS:
            fehler.append(f"Rätsel {r['id']}: unbekannter Status „{r['status']}“ (erlaubt: {', '.join(GUELTIGE_STATUS)})")
        if r.get("termin"):
            try:
                datetime.datetime.strptime(r["termin"], "%Y-%m-%d %H:%M")
            except (TypeError, ValueError):
                fehler.append(f"Rätsel {r['id']}: „termin“ muss so aussehen: 2026-10-25 18:30")
        if r.get("thema") and r["thema"] not in cfg.get("themen", {}):
            print(f"⚠ Rätsel {r['id']}: Thema „{r['thema']}“ ist nicht in config.json.")
        if r["kategorie"] not in cfg["kategorien"]:
            print(f"⚠ Rätsel {r['id']}: Kategorie „{r['kategorie']}“ ist nicht in config.json – nutze Standard-Design.")
        optionen = r.get("optionen")
        if optionen is not None:
            if not isinstance(optionen, list) or not 2 <= len(optionen) <= 4:
                fehler.append(f"Rätsel {r['id']}: „optionen“ braucht eine Liste mit 2 bis 4 Antworten")
            elif str(r["loesung"]) not in [str(o) for o in optionen]:
                print(f"⚠ Rätsel {r['id']}: Die Lösung „{r['loesung']}“ steht nicht wörtlich unter den Optionen.")
        if r.get("typ") == "finde":
            fehlend = [f for f in ("zeichen", "ziel", "spalten", "zeilen", "position") if f not in r]
            if fehlend:
                fehler.append(f"Rätsel {r['id']}: Für „finde“ fehlen: {', '.join(fehlend)}")
            elif not (1 <= r["position"][0] <= r["zeilen"] and 1 <= r["position"][1] <= r["spalten"]):
                fehler.append(f"Rätsel {r['id']}: „position“ liegt außerhalb des Gitters")
    if fehler:
        sys.exit("✗ Fehler in raetsel.json:\n  " + "\n  ".join(fehler))


def dateiname(r):
    return f"{int(r['id']):03d}_{r['kategorie']}"


# ---------------------------------------------------------------- Texte (Caption, Kommentar)

def caption_text(r, cfg):
    from renderer import einstellungen
    kat = einstellungen(cfg, r)
    zufall = random.Random(f"{r['id']}-{r['kategorie']}")  # gleiche ID → gleiche Caption
    captions = kat.get("captions") or ["Kannst du das lösen? 🧠"]
    from renderer import countdown_fuer
    caption = zufall.choice(captions).replace("{countdown}", str(round(countdown_fuer(cfg, r))))

    allgemein = list(cfg.get("hashtags_allgemein", []))
    anzahl = max(5, min(8, cfg.get("hashtags_anzahl", 7)))
    thema = [h for h in kat.get("thema_hashtags", []) if h not in allgemein][:3]  # Themen-Hashtags immer dabei
    eigene = [h for h in kat.get("hashtags", []) if h not in allgemein and h not in thema]
    zusatz = thema + zufall.sample(eigene, min(len(eigene), max(0, anzahl - len(allgemein) - len(thema))))
    hashtags = (allgemein[:1] + zusatz + allgemein[1:])[:anzahl]

    aufruf = kat.get("caption_aufruf") or cfg.get("caption_aufruf", "Schreib deine Antwort in die Kommentare 👇")
    if cfg.get("aufloesung", {}).get("zeigen"):
        aufruf = kat.get("caption_aufruf_aufloesung") or cfg.get("caption_aufruf_aufloesung") or aufruf
    elif r.get("optionen"):
        from renderer import buchstaben_text
        aufruf = kat.get("caption_aufruf_optionen") or cfg.get("caption_aufruf_optionen") or aufruf
        aufruf = aufruf.replace("{buchstaben}", buchstaben_text(len(r["optionen"])))
    return f"{caption}\n{aufruf}\n\n{' '.join(hashtags)}\n"


def kommentar_text(r, cfg):
    werte = dict(r)
    optionen = [str(o) for o in r.get("optionen") or []]
    if str(r["loesung"]) in optionen:  # bei A/B/C den Buchstaben mit angeben
        werte["loesung"] = f"{'ABCD'[optionen.index(str(r['loesung']))]} – {r['loesung']}"
    vorlage = cfg["kommentar_vorlage"]
    if cfg.get("aufloesung", {}).get("zeigen"):  # Lösung steht schon im Video → Kommentar fragt nach ✅/❌
        vorlage = cfg.get("kommentar_vorlage_aufloesung") or vorlage
    return vorlage.format(**werte).strip() + "\n"


# ---------------------------------------------------------------- Musik & Rendering

def musik_dateien(cfg):
    ordner = os.path.join(BASIS, cfg["musik"]["ordner"])
    if not os.path.isdir(ordner):
        return []
    return sorted(os.path.join(ordner, f) for f in os.listdir(ordner)
                  if f.lower().endswith(MUSIK_ENDUNGEN) and not f.startswith("."))


def soundeffekte_schreiben(renderer, ziel, rate=44100):
    """Erzeugt Countdown-Ticks und einen Gong bei „Zeit um!“ als WAV – ganz ohne Sound-Dateien."""
    puffer = array.array("f", bytes(4 * round(renderer.dauer * rate)))

    def ton(start, dauer, frequenzen, lautstaerke, abklingen):
        s0 = round(start * rate)
        for i in range(round(dauer * rate)):
            if s0 + i >= len(puffer):
                break
            t = i / rate
            huelle = math.exp(-t * abklingen) * min(1.0, t * 2000)  # kurzes Einschwingen gegen Knacken
            puffer[s0 + i] += lautstaerke * huelle * sum(math.sin(2 * math.pi * f * t) * g for f, g in frequenzen)

    for zeit, kritisch in renderer.ticks():
        if kritisch:  # höher, lauter und doppelt („Tick-Tack“) in den letzten Sekunden
            ton(zeit, 0.08, [(1760, 1.0), (3520, 0.3)], 0.45, 55)
            ton(zeit + 0.5, 0.05, [(1320, 1.0)], 0.2, 70)
        else:
            ton(zeit, 0.06, [(1200, 1.0), (2400, 0.25)], 0.28, 70)
    for zeit in renderer.pause_ticks():  # leises Ticken bei „Auflösung in 3 · 2 · 1“
        ton(zeit, 0.06, [(990, 1.0)], 0.2, 60)
    if renderer.aufloesung_zeitpunkt() is not None:  # heller Dur-Klang, wenn die Lösung erscheint
        ton(renderer.aufloesung_zeitpunkt(), 1.2, [(1046.5, 1.0), (1318.5, 0.7), (1568, 0.5)], 0.25, 4)
    if renderer.tipp_zeitpunkt() is not None:  # helles „Pling“, wenn der Tipp erscheint
        ton(renderer.tipp_zeitpunkt(), 0.6, [(1568, 1.0), (2349, 0.45)], 0.22, 8)
    if renderer.t_zeit_um > 0:
        ton(renderer.zeit_um_zeitpunkt(), 1.4, [(523.25, 1.0), (784, 0.6), (1046.5, 0.35)], 0.35, 3.2)

    daten = array.array("h", (max(-32767, min(32767, round(x * 32767))) for x in puffer))
    with wave.open(ziel, "wb") as w:
        w.setnchannels(1)
        w.setsampwidth(2)
        w.setframerate(rate)
        w.writeframes(daten.tobytes())


def musik_waehlen(cfg, ohne_musik):
    """Rückgabe (Pfad, generiert?): eigene Tracks haben Vorrang, sonst erzeugt das Skript die Musik selbst."""
    modus = cfg["musik"].get("modus", "auto")
    if ohne_musik or modus == "aus":
        return None, False
    dateien = musik_dateien(cfg) if modus in ("auto", "dateien") else []
    if dateien:
        return random.choice(dateien), False
    return (None, True) if modus in ("auto", "generiert") else (None, False)


def ffmpeg_befehl(ffmpeg, renderer, ziel, cfg, musik, effekte, musik_generiert=False):
    v = cfg["video"]
    dauer = renderer.dauer
    befehl = [ffmpeg, "-y", "-loglevel", "error",
              "-f", "rawvideo", "-pix_fmt", "rgb24", "-s", f"{renderer.breite}x{renderer.hoehe}",
              "-r", str(renderer.fps), "-i", "-"]
    filter_teile, spuren, eingang = [], [], 1
    if musik and musik_generiert:
        # Erzeugte Musik hat genau Videolänge und blendet selbst passend zum Loop aus
        befehl += ["-i", musik]
        filter_teile.append(f"[{eingang}:a]volume={cfg['musik']['lautstaerke']}[musik]")
        spuren.append("[musik]")
        eingang += 1
    elif musik:
        m = cfg["musik"]
        aus = min(m["ausblenden"], dauer)
        befehl += ["-stream_loop", "-1", "-i", musik]
        filter_teile.append(f"[{eingang}:a]atrim=0:{dauer:.3f},asetpts=N/SR/TB,"
                            f"afade=t=in:st=0:d={m['einblenden']},"
                            f"afade=t=out:st={dauer - aus:.3f}:d={aus},"
                            f"volume={m['lautstaerke']}[musik]")
        spuren.append("[musik]")
        eingang += 1
    if effekte:
        befehl += ["-i", effekte]
        filter_teile.append(f"[{eingang}:a]volume={cfg['soundeffekte']['lautstaerke']},"
                            f"aformat=channel_layouts=stereo,apad[sfx]")
        spuren.append("[sfx]")
        eingang += 1
    if spuren:
        if len(spuren) == 2:
            filter_teile.append("[musik][sfx]amix=inputs=2:duration=first:normalize=0[a]")
        else:
            filter_teile.append(f"{spuren[0]}anull[a]")
        ausgang = "[a]"
        lufs = cfg.get("audio", {}).get("lautheit_lufs")
        if lufs is not None:  # auf übliche Social-Media-Lautheit bringen
            filter_teile.append(f"[a]loudnorm=I={lufs}:TP=-1.5:LRA=11[laut]")
            ausgang = "[laut]"
        befehl += ["-filter_complex", ";".join(filter_teile), "-map", "0:v", "-map", ausgang]
    else:
        # Stille Tonspur, damit sich das Video überall problemlos weiterbearbeiten lässt
        befehl += ["-f", "lavfi", "-i", "anullsrc=r=44100:cl=stereo", "-map", "0:v", "-map", "1:a"]
    befehl += ["-t", f"{dauer:.3f}",
               "-c:v", "libx264", "-preset", v["preset"], "-crf", str(v["qualitaet_crf"]),
               "-profile:v", "high", "-pix_fmt", "yuv420p", "-r", str(renderer.fps),
               "-c:a", "aac", "-b:a", "192k", "-ar", "44100",
               "-movflags", "+faststart", ziel]
    return befehl


def erzeuge(r, cfg, ffmpeg, nur_vorschau=False, ohne_musik=False, ohne_effekte=False):
    from renderer import RaetselRenderer

    ausgabe = os.path.join(BASIS, cfg["ausgabe"]["ordner"])
    os.makedirs(ausgabe, exist_ok=True)
    name = dateiname(r)
    renderer = RaetselRenderer(cfg, r, BASIS)

    vorschau = os.path.join(ausgabe, f"{name}.png")
    renderer.frame(renderer.vorschau_zeitpunkt()).save(vorschau)
    if renderer.mit_aufloesung:
        renderer.frame(renderer.aufloesung_vorschau_zeitpunkt()).save(os.path.join(ausgabe, f"{name}_aufloesung.png"))

    if nur_vorschau:
        print(f"  ✓ Vorschau: {os.path.relpath(vorschau, BASIS)} (Schriftgröße {renderer.raetsel_schriftgroesse})")
        return

    musik, musik_generiert = musik_waehlen(cfg, ohne_musik)
    if musik_generiert:
        from musik import erzeuge_musik
        musik = os.path.join(ausgabe, f".{name}.musik.wav")
        erzeuge_musik(renderer, musik, saat=f"{r['id']}-{r['kategorie']}")
    video = os.path.join(ausgabe, f"{name}.mp4")
    tmp = os.path.join(ausgabe, f".{name}.tmp.mp4")
    effekte = None
    if cfg.get("soundeffekte", {}).get("aktiv") and not ohne_effekte:
        effekte = os.path.join(ausgabe, f".{name}.sfx.wav")
        soundeffekte_schreiben(renderer, effekte)

    start = time.time()
    try:
        prozess = subprocess.Popen(ffmpeg_befehl(ffmpeg, renderer, tmp, cfg, musik, effekte, musik_generiert),
                                   stdin=subprocess.PIPE)
        try:
            for i in range(renderer.frame_anzahl):
                prozess.stdin.write(renderer.frame(i / renderer.fps).tobytes())
            prozess.stdin.close()
        except BrokenPipeError:
            pass
        if prozess.wait() != 0:
            if os.path.exists(tmp):
                os.remove(tmp)
            raise RuntimeError("FFmpeg ist mit einem Fehler abgebrochen (siehe Meldung oben).")
    finally:
        for hilfsdatei in (effekte, musik if musik_generiert else None):
            if hilfsdatei and os.path.exists(hilfsdatei):
                os.remove(hilfsdatei)
    os.replace(tmp, video)

    with open(os.path.join(ausgabe, f"{name}_caption.txt"), "w", encoding="utf-8") as f:
        f.write(caption_text(r, cfg))
    kommentar = kommentar_text(r, cfg)
    with open(os.path.join(ausgabe, f"{name}_kommentar.txt"), "w", encoding="utf-8") as f:
        f.write(kommentar)

    musik_info = "erzeugte Musik" if musik_generiert else os.path.basename(musik) if musik else "ohne Musik"
    musik_info += ", mit Soundeffekten" if effekte else ""
    print(f"  ✓ {os.path.relpath(video, BASIS)}  ({renderer.dauer:.1f} s, {musik_info}, "
          f"Schriftgröße {renderer.raetsel_schriftgroesse}, {time.time() - start:.1f} s Renderzeit)")
    laenge = len(kommentar.strip().encode("utf-16-le")) // 2  # Emojis zählen hier doppelt – wie bei TikTok
    if laenge > KOMMENTAR_MAX:
        print(f"  ⚠ Auflösungs-Kommentar hat {laenge} Zeichen – TikTok erlaubt {KOMMENTAR_MAX}. Erklärung kürzen!")


# ---------------------------------------------------------------- Posting-Plan

WOCHENTAGE = ["Mo", "Di", "Mi", "Do", "Fr", "Sa", "So"]


def youtube_titel(caption):
    """Erste Caption-Zeile + #shorts – YouTube erlaubt höchstens 100 Zeichen und kein < oder >."""
    titel = caption.strip().split("\n")[0].replace("<", "").replace(">", "").strip()
    return titel + " #shorts" if len(titel) + 8 < 100 else titel[:99]


def plan_als_excel(pfad, zeilen):
    try:
        import openpyxl
    except ImportError:
        pip_installieren()
        import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    mappe = openpyxl.Workbook()
    blatt = mappe.active
    blatt.title = "Posting-Plan"
    spalten = [("Datum", 12), ("Tag", 5), ("Uhrzeit", 8), ("Video", 22), ("Schwierigkeit", 13),
               ("Thema", 12), ("YouTube-Titel", 45), ("Caption", 70), ("Lösung", 18), ("Eingeplant ✓", 13)]
    blatt.append([name for name, _ in spalten])
    for i, (_, breite) in enumerate(spalten, start=1):
        blatt.column_dimensions[openpyxl.utils.get_column_letter(i)].width = breite
    for zelle in blatt[1]:
        zelle.font = Font(bold=True, color="FFFFFF")
        zelle.fill = PatternFill("solid", fgColor="2D1B2E")
    for zeile in zeilen:
        blatt.append(zeile)
    for zeile in blatt.iter_rows(min_row=2):
        zeile[0].number_format = "DD.MM.YYYY"
        for zelle in zeile:
            zelle.alignment = Alignment(vertical="top", wrap_text=True)
    blatt.freeze_panes = "A2"
    blatt.auto_filter.ref = blatt.dimensions
    mappe.save(pfad)


def posting_plan(raetsel, cfg, tage, uhrzeiten, start):
    """Verteilt fertige, noch nicht gepostete Videos auf Tage und Uhrzeiten – als Vorlage zum Einplanen.
    Jedes Rätsel bekommt seinen Termin; ist der vorbei, zählt es automatisch als gepostet."""
    ausgabe = os.path.join(BASIS, cfg["ausgabe"]["ordner"])
    from renderer import im_zeitraum
    # Termine ab dem Starttag werden neu vergeben – so kann man denselben Plan gefahrlos neu erzeugen
    for r in raetsel:
        if r.get("termin") and r["status"] != "gepostet" and r["termin"][:10] >= start.isoformat():
            del r["termin"]
    fertig = [r for r in raetsel if r["status"] != "gepostet" and not r.get("termin")
              and os.path.exists(os.path.join(ausgabe, dateiname(r) + ".mp4"))]
    themen = cfg.get("themen", {})
    zeilen = [f"POSTING-PLAN IQKnacker – {len(uhrzeiten)} Videos pro Tag", "=" * 44, ""]
    tabelle = []
    i = 0
    for tag in range(tage):
        datum = start + datetime.timedelta(days=tag)
        for uhrzeit in uhrzeiten:
            erlaubt = [r for r in fertig if im_zeitraum(cfg, r, datum)]
            if not erlaubt:
                break
            # Saison-Rätsel zuerst – sie laufen sonst ab; abwechselnd mit normalen, damit es nicht eintönig wird
            saison = [r for r in erlaubt if themen.get(r.get("thema") or "", {}).get("zeitraum")]
            r = saison[0] if saison and (i % 2 == 0 or len(saison) >= len(erlaubt)) else \
                next((x for x in erlaubt if x not in saison), erlaubt[0])
            fertig.remove(r)
            r["termin"] = f"{datum.isoformat()} {uhrzeit}"
            caption = caption_text(r, cfg).strip()
            tabelle.append([datum, WOCHENTAGE[datum.weekday()], uhrzeit, f"{dateiname(r)}.mp4", r["schwierigkeit"],
                            r.get("thema") or "", youtube_titel(caption), caption, str(r["loesung"]), ""])
            i += 1
            zeilen += [f"{WOCHENTAGE[datum.weekday()]}, {datum:%d.%m.} um {uhrzeit} Uhr  →  {dateiname(r)}.mp4"
                       f"  ({r['schwierigkeit']}{', ' + r['thema'] if r.get('thema') else ''})",
                       "", "Caption:", caption, "", "YouTube-Titel:", youtube_titel(caption), "",
                       "Kommentar zum Anpinnen:", kommentar_text(r, cfg).strip(), "", "-" * 44, ""]
    if i < tage * len(uhrzeiten):
        zeilen.append(f"Hinweis: Nur {i} fertige Videos vorhanden – vorher python generate.py ausführen.")
    pfad = os.path.join(ausgabe, "postingplan.txt")
    with open(pfad, "w", encoding="utf-8") as f:
        f.write("\n".join(zeilen) + "\n")
    excel = os.path.join(ausgabe, "postingplan.xlsx")
    plan_als_excel(excel, tabelle)
    speichere_raetsel(raetsel)
    print(f"✓ {os.path.relpath(pfad, BASIS)} und {os.path.basename(excel)}: "
          f"{i} Videos auf {tage} Tage verteilt ({', '.join(uhrzeiten)} Uhr)")


# ---------------------------------------------------------------- Kommandozeile

def main():
    parser = argparse.ArgumentParser(description="IQKnacker: erzeugt Rätsel-Videos aus raetsel.json")
    parser.add_argument("--id", type=int, nargs="+", help="nur diese Rätsel-ID(s) erzeugen (egal welcher Status)")
    parser.add_argument("--alle", action="store_true", help="alle Rätsel erzeugen, egal welcher Status")
    parser.add_argument("--vorschau", action="store_true", help="nur Vorschaubilder, kein Video, Status bleibt")
    parser.add_argument("--ohne-musik", action="store_true", help="Video ohne Musik erzeugen")
    parser.add_argument("--ohne-effekte", action="store_true", help="ohne Ticken und Gong")
    parser.add_argument("--ohne-aufloesung", action="store_true",
                        help="Lösung nicht im Video zeigen (Endkarte wie früher, z. B. zum Vergleichen)")
    parser.add_argument("--liste", action="store_true", help="alle Rätsel mit Status anzeigen")
    parser.add_argument("--gepostet", type=int, nargs="+", metavar="ID",
                        help="diese Rätsel-ID(s) als gepostet markieren")
    parser.add_argument("--check", action="store_true", help="Installation prüfen")
    parser.add_argument("--neu-emoji", type=int, metavar="ANZAHL", help="neue Emoji-Gleichungen anlegen")
    parser.add_argument("--neu-finde", type=int, metavar="ANZAHL", help="neue „Finde die 9“-Rätsel anlegen")
    parser.add_argument("--plan", type=int, metavar="TAGE", help="Posting-Plan für so viele Tage erstellen")
    parser.add_argument("--uhrzeiten", default="12:30,18:30,21:00",
                        help="Uhrzeiten für --plan, Komma-getrennt (Standard: 12:30,18:30,21:00)")
    parser.add_argument("--start", help="erster Tag für --plan als JJJJ-MM-TT (Standard: morgen)")
    parser.add_argument("--schwierigkeit", choices=["leicht", "mittel", "schwer"],
                        help="für --neu-emoji/--neu-finde: nur diese Stufe (sonst gemischt)")
    args = parser.parse_args()

    if args.check:
        check()
        return

    cfg = lade_json(CONFIG_DATEI)
    if args.ohne_aufloesung:
        cfg.setdefault("aufloesung", {})["zeigen"] = False
    raetsel = lade_json(RAETSEL_DATEI)
    pruefe_raetsel(raetsel, cfg)
    if termine_abhaken(raetsel):
        speichere_raetsel(raetsel)

    if args.neu_emoji or args.neu_finde:
        import automatik
        neu = []
        if args.neu_emoji:
            neu += automatik.erzeuge("emoji", args.neu_emoji, raetsel, args.schwierigkeit)
        if args.neu_finde:
            neu += automatik.erzeuge("finde", args.neu_finde, raetsel, args.schwierigkeit)
        speichere_raetsel(raetsel)
        for r in neu:
            frage = r["frage"].replace("\n", "  |  ").replace("*", "")
            print(f"+ {r['id']:>3} {r['kategorie']:<6} {r['schwierigkeit']:<7} {frage}  →  {r['loesung']}")
        print(f"{len(neu)} Rätsel angelegt. Videos erzeugen mit: python generate.py")
        return

    if args.gepostet:
        treffer = [r for r in raetsel if r["id"] in args.gepostet]
        fehlend = set(args.gepostet) - {r["id"] for r in treffer}
        if fehlend:
            sys.exit(f"✗ Rätsel-ID(s) nicht gefunden: {', '.join(map(str, sorted(fehlend)))}")
        for r in treffer:
            r["status"] = "gepostet"
            r.pop("termin", None)
        speichere_raetsel(raetsel)
        print(f"✓ Als gepostet markiert: {', '.join(str(r['id']) for r in treffer)}")
        return

    if args.plan:
        start = datetime.date.fromisoformat(args.start) if args.start else \
            datetime.date.today() + datetime.timedelta(days=1)
        posting_plan(raetsel, cfg, args.plan, [u.strip() for u in args.uhrzeiten.split(",") if u.strip()], start)
        return

    if args.liste:
        for r in raetsel:
            frage = r["frage"].replace("\n", " ").replace("*", "")
            thema = f"[{r['thema']}] " if r.get("thema") else ""
            termin = f"({r['termin']}) " if r.get("termin") and r["status"] != "gepostet" else ""
            print(f"{r['id']:>4}  {r['status']:<9} {r['kategorie']:<7} {termin}{thema}{frage[:60]}")
        return

    if args.id:
        auswahl = [r for r in raetsel if r["id"] in args.id]
        fehlend = set(args.id) - {r["id"] for r in auswahl}
        if fehlend:
            sys.exit(f"✗ Rätsel-ID(s) nicht gefunden: {', '.join(map(str, sorted(fehlend)))}")
    elif args.alle:
        auswahl = raetsel
    else:
        auswahl = [r for r in raetsel if r["status"] == "neu"]

    if not auswahl:
        print("Keine neuen Rätsel. Neue Einträge mit \"status\": \"neu\" in raetsel.json anlegen "
              "oder --id / --alle nutzen.")
        return

    pakete_sicherstellen()
    ffmpeg = None if args.vorschau else ffmpeg_finden()

    print(f"Erzeuge {len(auswahl)} {'Vorschau(en)' if args.vorschau else 'Video(s)'} ...")
    fehler = 0
    for r in auswahl:
        print(f"• Rätsel {r['id']} ({r['kategorie']})")
        try:
            erzeuge(r, cfg, ffmpeg, nur_vorschau=args.vorschau, ohne_musik=args.ohne_musik,
                    ohne_effekte=args.ohne_effekte)
        except Exception as e:
            fehler += 1
            print(f"  ✗ Fehler: {e}")
            continue
        if not args.vorschau and r["status"] == "neu":
            r["status"] = "erstellt"
            speichere_raetsel(raetsel)  # nach jedem Video speichern, damit nichts verloren geht

    print("Fertig." if not fehler else f"Fertig mit {fehler} Fehler(n).")
    if fehler:
        sys.exit(1)


if __name__ == "__main__":
    main()
